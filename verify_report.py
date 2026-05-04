"""Verify the expanded report structure and estimate page count."""
import docx

doc = docx.Document(r"Report\PROJECT REPORT FINAL VERSION v2 - EXPANDED.docx")

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Count images
image_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        image_count += 1
print(f"Total images: {image_count}")

# Count total characters (rough page estimate)
total_chars = sum(len(p.text) for p in doc.paragraphs)
print(f"Total characters: {total_chars}")

# Rough estimate: ~2500 chars per page for academic report with figures
est_pages = total_chars / 2500
print(f"Estimated pages (text only): {est_pages:.1f}")
print(f"Estimated pages (with figures/tables/spacing): {est_pages * 1.3:.1f}")

# Check structure preserved
print("\n=== STRUCTURE CHECK ===")
headings = []
for p in doc.paragraphs:
    text = p.text.strip()
    if text.startswith("CHAPTER ") or text.startswith("APPENDIX"):
        headings.append(text)
    elif text.startswith("TABLE OF CONTENTS"):
        headings.append(text)
    elif text.startswith("LIST OF"):
        headings.append(text)
    elif text.startswith("ABBREVIATIONS"):
        headings.append(text)
    elif text.startswith("REFERENCES"):
        headings.append(text)

for h in headings:
    print(f"  {h}")

# Count non-empty paragraphs with text content
body_paras = [p for p in doc.paragraphs if len(p.text.strip()) > 50]
print(f"\nSubstantial paragraphs (>50 chars): {len(body_paras)}")

# Compare with original
doc_orig = docx.Document(r"Report\PROJECT REPORT FINAL VERSION v2.docx")
orig_chars = sum(len(p.text) for p in doc_orig.paragraphs)
print(f"\nOriginal characters: {orig_chars}")
print(f"Expanded characters: {total_chars}")
print(f"Character increase: {total_chars - orig_chars} ({(total_chars - orig_chars)/orig_chars*100:.1f}%)")
orig_est = orig_chars / 2500
print(f"Original est. pages: {orig_est:.1f}")
print(f"Expansion ratio: {total_chars/orig_chars:.2f}x")
