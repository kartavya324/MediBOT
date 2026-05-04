"""Extract report structure and content for analysis."""
import docx
import json
import os

doc = docx.Document(r"Report\PROJECT REPORT FINAL VERSION v2.docx")

# Extract all paragraphs with their styles
output = []
for i, para in enumerate(doc.paragraphs):
    style_name = para.style.name if para.style else "None"
    text = para.text.strip()
    if text or style_name.startswith("Heading"):
        output.append({
            "idx": i,
            "style": style_name,
            "text": text[:200] if len(text) > 200 else text,
            "full_len": len(text),
            "bold": any(run.bold for run in para.runs if run.bold),
            "alignment": str(para.alignment),
        })

# Save to file
with open("report_structure.json", "w", encoding="utf-8") as f:
    json.dump(output, f, indent=2, ensure_ascii=False)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Non-empty paragraphs: {len(output)}")

# Count tables
print(f"Total tables: {len(doc.tables)}")

# Count images
image_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.reltype:
        image_count += 1
print(f"Total images: {image_count}")

# Print sections/headings
print("\n=== HEADINGS/STRUCTURE ===")
for item in output:
    if item["style"].startswith("Heading") or item["bold"]:
        print(f"  [{item['style']}] {item['text']}")

# Print all content for detailed analysis
print("\n=== FULL CONTENT DUMP ===")
for para in doc.paragraphs:
    text = para.text.strip()
    style = para.style.name if para.style else "None"
    if text:
        print(f"[{style}] {text}")
